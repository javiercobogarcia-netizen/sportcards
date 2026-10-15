import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from datetime import datetime

DATA_DIR = os.getenv('DATA_DIR', os.path.dirname(__file__))
EXPORTS_DIR = os.path.join(DATA_DIR, 'exports')

def export_collection(cards):
    os.makedirs(EXPORTS_DIR, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Colección'

    header_fill = PatternFill(start_color='CC0000', end_color='CC0000', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    owned_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    sold_fill = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')

    headers = [
        'ID', 'Jugador', 'Año', 'Deporte', 'Equipo', 'Set/Colección',
        'Variante', 'Nº Carta', 'Numeración', 'Gradada', 'Empresa Gradeo',
        'Nota', 'Precio Compra (€)', 'Fecha Compra', 'Precio Mercado (€)',
        'Precio Venta (€)', 'Fecha Venta', 'Beneficio (€)', 'ROI (%)', 'Estado', 'Notas'
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[1].height = 25

    col_widths = [5,22,6,12,18,25,15,8,10,8,12,8,16,14,16,16,12,14,10,8,25]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row_idx, card in enumerate(cards, 2):
        purchase = card.get('purchase_price') or 0
        sale = card.get('sale_price') or 0
        market = card.get('market_price') or 0

        if card.get('status') == 'sold' and purchase and sale:
            profit = round(sale - purchase, 2)
            roi = round((profit / purchase) * 100, 1) if purchase else 0
        elif card.get('status') == 'owned' and market:
            profit = round(market - purchase, 2)
            roi = round((profit / purchase) * 100, 1) if purchase else 0
        else:
            profit = None
            roi = None

        values = [
            card.get('id'),
            card.get('player'),
            card.get('year'),
            card.get('sport'),
            card.get('team'),
            card.get('set_name'),
            card.get('variant'),
            card.get('card_number'),
            card.get('numbering'),
            'Sí' if card.get('graded') else 'No',
            card.get('grading_company'),
            card.get('grade'),
            purchase if purchase else None,
            card.get('purchase_date'),
            market if market else None,
            sale if sale else None,
            card.get('sale_date'),
            profit,
            roi,
            card.get('status', 'owned'),
            card.get('notes')
        ]

        fill = sold_fill if card.get('status') == 'sold' else owned_fill

        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill = fill
            cell.alignment = Alignment(vertical='center')

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet('Resumen')
    owned = [c for c in cards if c.get('status') == 'owned']
    sold = [c for c in cards if c.get('status') == 'sold']
    total_invested = sum(c.get('purchase_price') or 0 for c in owned)
    current_value = sum(c.get('market_price') or c.get('purchase_price') or 0 for c in owned)
    total_profit_sold = sum((c.get('sale_price') or 0) - (c.get('purchase_price') or 0) for c in sold)

    summary_data = [
        ('RESUMEN DE COLECCIÓN', ''),
        ('Fecha exportación', datetime.now().strftime('%d/%m/%Y %H:%M')),
        ('', ''),
        ('Total cartas', len(cards)),
        ('Cartas en posesión', len(owned)),
        ('Cartas vendidas', len(sold)),
        ('', ''),
        ('Total invertido (cartas actuales)', f'{total_invested:.2f} €'),
        ('Valor actual de mercado', f'{current_value:.2f} €'),
        ('Beneficio/pérdida latente', f'{current_value - total_invested:.2f} €'),
        ('Beneficio realizado (vendidas)', f'{total_profit_sold:.2f} €'),
    ]

    for row_idx, (label, value) in enumerate(summary_data, 1):
        ws2.cell(row=row_idx, column=1, value=label).font = Font(bold=True if row_idx == 1 else False)
        ws2.cell(row=row_idx, column=2, value=value)

    ws2.column_dimensions['A'].width = 35
    ws2.column_dimensions['B'].width = 20

    filename = f'coleccion_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    path = os.path.join(EXPORTS_DIR, filename)
    wb.save(path)
    return path, filename
